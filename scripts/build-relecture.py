"""Tableur de relecture du contenu V5 (bloc 2.6), à partir de content/v5/relecture.json.
   python3 scripts/build-relecture.py   (après node scripts/build-content.mjs)
Produit content/v5/relecture-bloc2.xlsx. Valentin remplit les colonnes jaunes, puis le renvoie :
scripts/build-validation.py génère le SQL qui valide ce qui est marqué « oui »."""
import json
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
data = json.loads((ROOT / 'content/v5/relecture.json').read_text())
OUT = ROOT / 'content/v5/relecture-bloc2.xlsx'

F = 'Arial'
HEAD = PatternFill('solid', fgColor='1F3A5F')
INPUT = PatternFill('solid', fgColor='FFF2A8')
EXAMPLE = PatternFill('solid', fgColor='E8F0E0')
thin = Side(style='thin', color='D0D0D0')
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
WRAP = Alignment(wrap_text=True, vertical='top')

wb = Workbook()


def table(ws, headers, rows, widths, inputs):
    for c, (h, w) in enumerate(zip(headers, widths), start=1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.font = Font(name=F, bold=True, color='FFFFFF')
        cell.fill = INPUT if h in inputs else HEAD
        if h in inputs:
            cell.font = Font(name=F, bold=True, color='000000')
        cell.alignment = Alignment(wrap_text=True, vertical='center')
        cell.border = BORDER
        ws.column_dimensions[get_column_letter(c)].width = w
    for r, row in enumerate(rows, start=2):
        for c, v in enumerate(row, start=1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.font = Font(name=F, size=9)
            cell.alignment = WRAP
            cell.border = BORDER
            if headers[c - 1] in inputs:
                cell.fill = INPUT
    ws.freeze_panes = 'C2'
    ws.auto_filter.ref = f'A1:{get_column_letter(len(headers))}{len(rows) + 1}'
    ws.row_dimensions[1].height = 32
    return len(rows) + 1


def yesno(ws, col, last):
    dv = DataValidation(type='list', formula1='"oui,non,à revoir"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f'{col}2:{col}{last}')


# ------------------------------------------------------------------ Exercices
ws = wb.active
ws.title = 'Exercices'
heads = ['Identifiant', 'Nom', 'Statut', 'Validé', 'Vidéo OK', 'Commentaire', 'Sources', 'Discipline', 'Zones',
         'Objectifs', 'Intensité (1 à 5)', 'Sauts', 'Matériel nécessaire', 'Par défaut', 'Niveaux', 'Muscles principaux',
         'Muscles secondaires', 'MET (Compendium 2024)', 'Consignes', 'Précaution', 'Respiration',
         'Erreurs fréquentes', 'Vidéo proposée']
rows = []
for e in data['exercices']:
    v = e['video']
    rows.append([e['id'], e['nom'], e['statut'], None, None, None,
                 '\n'.join(f"{s['label']} : {s['url']}" for s in e['sources']),
                 e['discipline'], e['zones'], e['objectifs'],
                 e['intensite'], e['sauts'], e['materiel'], e['mode'], e['niveaux'], e['muscles_principaux'],
                 e['muscles_secondaires'], e['met'], e['consignes'], e['precaution'], e['respiration'], e['erreurs'],
                 f"{v['title']} ({'français' if v['lang'] == 'fr' else 'anglais'})\n{v['url']}" if v else ''])
last_ex = table(ws, heads, rows,
                [20, 30, 14, 10, 10, 30, 60, 13, 18, 18, 9, 7, 14, 13, 40, 28, 28, 30, 55, 35, 30, 38, 50],
                {'Validé', 'Vidéo OK', 'Commentaire'})
yesno(ws, 'D', last_ex)
yesno(ws, 'E', last_ex)
for r in range(2, last_ex + 1):
    url = ws.cell(row=r, column=23).value.split('\n')[-1] if ws.cell(row=r, column=23).value else None
    if url:
        ws.cell(row=r, column=23).hyperlink = url
        ws.cell(row=r, column=23).font = Font(name=F, size=9, color='0563C1', underline='single')

# ------------------------------------------------------------------ Entraînements
ws2 = wb.create_sheet('Entraînements')
kinds = {'programme': 'Jour de programme', 'rapide': 'Démarrage rapide', 'decouverte': 'Découverte', 'express': 'Express (hors ligne)'}
rows = [[t['id'], t['titre'], kinds[t['type']], None, None, t['niveau'], t['zones'], t['minutes'],
         round(t['metMinutes'] * 60 / 60, 0), t['tours'], t['pause'], 'oui' if t['impact'] else 'non',
         ', '.join(t['equipment']) or 'aucun', t['description'], t['contenu'],
         '\n'.join(f"{x['label']} : {x['url']}" for x in t['reperes']), '\n'.join(f"{x['label']} : {x['url']}" for x in t['sources_exercices'])]
        for t in data['entrainements']]
last_tr = table(ws2, ['Identifiant', 'Titre', 'Type', 'Validé', 'Commentaire', 'Niveau', 'Zones', 'Durée (min)',
                      'Calories pour 60 kg (estimation)', 'Tours', 'Pause (s)', 'Sauts', 'Matériel', 'Description',
                      'Contenu', 'Repères de construction', 'Sources des exercices'], rows,
                [26, 34, 18, 10, 30, 7, 20, 9, 12, 7, 8, 7, 10, 45, 60, 55, 70], {'Validé', 'Commentaire'})
yesno(ws2, 'D', last_tr)
# Calories : MET-minutes x poids / 60 (cadrage V5, section 8). Poids réglable dans « Mode d'emploi ».
WEIGHT = "'Mode d''emploi'!$B$21"
ws2.cell(row=1, column=9, value='Calories (estimation, poids du mode d\'emploi)')
for r, t in enumerate(data['entrainements'], start=2):
    ws2.cell(row=r, column=9, value=f'=ROUND({t["metMinutes"]}*{WEIGHT}/60,0)')

# ------------------------------------------------------------------ Programmes
ws3 = wb.create_sheet('Programmes')
rows = []
for p in data['programmes']:
    for j in p['jours']:
        rows.append([p['titre'], j['jour'], (j['jour'] - 1) // 7 + 1, j['entrainement'], j['titre'], j['minutes']])
last_pr = table(ws3, ['Programme', 'Jour', 'Semaine', 'Entraînement', 'Titre', 'Durée (min)'], rows,
                [24, 7, 9, 22, 38, 40], set())
ws3.freeze_panes = 'A2'
c = len(rows) + 3
ws3.cell(row=c, column=1, value='Programme').font = Font(name=F, bold=True)
ws3.cell(row=c, column=2, value='Validé').font = Font(name=F, bold=True)
ws3.cell(row=c, column=3, value='Commentaire').font = Font(name=F, bold=True)
ws3.cell(row=c, column=4, value='Description').font = Font(name=F, bold=True)
ws3.cell(row=c, column=5, value='Repères de construction').font = Font(name=F, bold=True)
ws3.cell(row=c, column=6, value='Sources des exercices').font = Font(name=F, bold=True)
ws3.column_dimensions['G'].width = 10
dv = DataValidation(type='list', formula1='"oui,non,à revoir"', allow_blank=True)
ws3.add_data_validation(dv)
for i, p in enumerate(data['programmes'], start=1):
    ws3.cell(row=c + i, column=1, value=p['titre']).font = Font(name=F, size=9)
    for col in (2, 3):
        ws3.cell(row=c + i, column=col).fill = INPUT
    dv.add(ws3.cell(row=c + i, column=2))
    d = ws3.cell(row=c + i, column=4, value=p['description'])
    d.font = Font(name=F, size=9)
    d.alignment = WRAP
    for col, val in ((5, '\n'.join(f"{x['label']} : {x['url']}" for x in p['reperes'])),
                     (6, f"{len(p['sources_exercices'])} sources : voir l'onglet Sources, colonne Programmes")):
        x = ws3.cell(row=c + i, column=col, value=val)
        x.font = Font(name=F, size=9)
        x.alignment = WRAP

# ------------------------------------------------------------------ Sources
ws4 = wb.create_sheet('Sources')
prog_urls = {p['titre']: {x['url'] for x in p['sources_exercices']} | {x['url'] for x in p['reperes']} for p in data['programmes']}
rows = []
for src in data['sources']:
    progs = [t for t, urls in prog_urls.items() if src['url'] in urls]
    n_tr = sum(1 for t in data['entrainements']
               if src['url'] in {x['url'] for x in t['reperes'] + t['sources_exercices']})
    rows.append([src['type'], src['titre'], src['url'], src['usages'], None,
                 ', '.join(src['exercices']), n_tr, ', '.join(progs)])
last_src = table(ws4, ['Type', 'Source', 'Lien', 'À quoi elle sert', "Nombre d'exercices", 'Exercices',
                       "Nombre d'entraînements", 'Programmes'], rows,
                 [22, 50, 45, 38, 11, 60, 13, 34], set())
ws4.freeze_panes = 'C2'
for r in range(2, last_src + 1):
    # Nombre d'exercices qui citent cette source : recherche du lien dans la colonne Sources de l'onglet Exercices
    ws4.cell(row=r, column=5, value=f'=SUMPRODUCT(--ISNUMBER(SEARCH(C{r},Exercices!$G$2:$G${last_ex})))')
    link = ws4.cell(row=r, column=3)
    link.hyperlink = link.value
    link.font = Font(name=F, size=9, color='0563C1', underline='single')

# ------------------------------------------------------------------ Mode d'emploi (en premier)
ws0 = wb.create_sheet("Mode d'emploi", 0)
ws0.column_dimensions['A'].width = 46
ws0.column_dimensions['B'].width = 16
ws0.column_dimensions['C'].width = 70
lines = [
    ('Relecture du contenu V5 (bloc 2)', None, None),
    ('Rien de ce tableur n\'apparaît dans l\'app tant que ce n\'est pas validé : tout est importé « à valider ».', None, None),
    (None, None, None),
    ('Ce que tu remplis (cases jaunes)', None, None),
    ('Validé', 'oui / non / à revoir', 'oui : l\'exercice ou l\'entraînement devient visible. non : il reste caché. à revoir : précise dans Commentaire.'),
    ('Vidéo OK', 'oui / non', 'non : je cherche une autre vidéo.'),
    ('Commentaire', 'texte libre', 'Ce qu\'il faut changer (consigne, niveau, intensité, vidéo...).'),
    ('Exemple', 'Validé = oui', 'Vidéo OK = non ; Commentaire = « vidéo trop rapide, en trouver une plus pédagogique »'),
    (None, None, None),
    ('Où en est la relecture', 'Nombre', None),
]
for r, (a, b, cc) in enumerate(lines, start=1):
    for col, v in ((1, a), (2, b), (3, cc)):
        cell = ws0.cell(row=r, column=col, value=v)
        cell.font = Font(name=F, size=10)
        cell.alignment = WRAP
ws0['A1'].font = Font(name=F, size=14, bold=True)
for r in (4, 10):
    ws0.cell(row=r, column=1).font = Font(name=F, bold=True)
for r in range(5, 8):
    ws0.cell(row=r, column=1).fill = INPUT
for col in (1, 2, 3):
    ws0.cell(row=8, column=col).fill = EXAMPLE
counts = [
    ('Exercices au total', f"=COUNTA(Exercices!A2:A{last_ex})"),
    ('Dont nouveaux', f'=COUNTIF(Exercices!C2:C{last_ex},"Nouveau")'),
    ('Exercices validés', f'=COUNTIF(Exercices!D2:D{last_ex},"oui")'),
    ('Exercices à revoir', f'=COUNTIF(Exercices!D2:D{last_ex},"à revoir")'),
    ('Vidéos refusées', f'=COUNTIF(Exercices!E2:E{last_ex},"non")'),
    ('Entraînements au total', f"=COUNTA('Entraînements'!A2:A{last_tr})"),
    ('Entraînements validés', f'=COUNTIF(\'Entraînements\'!D2:D{last_tr},"oui")'),
    ('Exercices restant à relire', f'=COUNTA(Exercices!A2:A{last_ex})-COUNTA(Exercices!D2:D{last_ex})'),
    ('Sources consultées', f'=COUNTA(Sources!C2:C{last_src})'),
]
for i, (lab, f) in enumerate(counts, start=11):
    ws0.cell(row=i, column=1, value=lab).font = Font(name=F, size=10)
    ws0.cell(row=i, column=2, value=f).font = Font(name=F, size=10, bold=True)
ws0.cell(row=21, column=1, value='Poids pour estimer les calories (kg)').font = Font(name=F, size=10)
ws0.cell(row=21, column=2, value=60).font = Font(name=F, size=10, color='0000FF')
ws0.cell(row=21, column=2).fill = INPUT
ws0.cell(row=21, column=3, value='Exemple : 60 kg. Change-le pour voir les calories des entraînements à ton poids.').font = Font(name=F, size=10)
n = 23
notes = [
    'À savoir',
    'Un entraînement validé ne s\'affiche que si tous ses exercices sont validés aussi.',
    'Les 43 exercices existants gardent leur nom, leurs consignes et leurs niveaux : seules les étiquettes, le tutoriel, la source et la vidéo sont nouveaux.',
    'Calories : MET x poids x durée (Compendium 2024). La colonne des entraînements prend 60 kg pour exemple ; l\'app utilisera le poids du profil.',
    'Sources : l\'onglet Sources liste chacune des sources consultées, avec son lien, à quoi elle a servi et qui la cite. Elles restent dans la base et dans ce tableur ; l\'app ne les affiche pas. Aucun texte n\'est copié.',
    'Vidéos : liens YouTube trouvés en ligne, en français quand c\'était possible. Je ne peux pas les regarder : c\'est à toi de vérifier qu\'elles montrent le bon mouvement.',
]
for i, t in enumerate(notes):
    cell = ws0.cell(row=n + i, column=1, value=t)
    cell.font = Font(name=F, size=10, bold=(i == 0))
    ws0.merge_cells(start_row=n + i, start_column=1, end_row=n + i, end_column=3)
    cell.alignment = WRAP
    ws0.row_dimensions[n + i].height = 30 if i else 18

wb.save(OUT)
print(OUT)
