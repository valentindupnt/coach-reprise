"""Après relecture : génère supabase/content-v5-validation.sql à partir du tableur rempli.
   python3 scripts/build-validation.py chemin/vers/relecture-bloc2.xlsx
- Exercice marqué « oui » : validé (approved = true).
- Entraînement marqué « oui » : validé seulement si tous ses exercices sont validés (dans le tableur ou déjà en base).
- Programme marqué « oui » : validé seulement si tous ses entraînements le sont.
Rien n'est jamais « dé-validé » par ce script : un « non » laisse l'élément tel qu'il est en base."""
import json
import sys
from pathlib import Path
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
src = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / 'content/v5/relecture-bloc2.xlsx'
wb = load_workbook(src, data_only=True)
data = json.loads((ROOT / 'content/v5/relecture.json').read_text())


def answers(ws, col):
    out = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0] and str(row[col] or '').strip().lower() == 'oui':
            out[row[0]] = True
    return out


ex_ok = answers(wb['Exercices'], 3)
tr_ok = answers(wb['Entraînements'], 3)
ws = wb['Programmes']
pr_ok = {}
for row in ws.iter_rows(values_only=True):
    if row[0] and str(row[1] or '').strip().lower() == 'oui' and not isinstance(row[1], int):
        pr_ok[row[0]] = True

q = lambda s: "'" + s.replace("'", "''") + "'"
lines = ['-- Validation du contenu V5 (bloc 2), générée par scripts/build-validation.py depuis ' + src.name,
         '-- À lancer APRÈS content-v5.sql. Relançable.', '']
if ex_ok:
    lines.append('update public.exercises set approved = true where id in (' + ', '.join(q(i) for i in sorted(ex_ok)) + ');')
if tr_ok:
    lines.append('-- Entraînements : seulement si tous leurs exercices sont validés')
    lines.append('update public.trainings t set approved = true where t.id in (' + ', '.join(q(i) for i in sorted(tr_ok)) + ')\n'
                 '  and not exists (select 1 from public.training_items ti join public.exercises e on e.id = ti.exercise_id\n'
                 '                  where ti.training_id = t.id and not e.approved);')
titles = {p['titre']: p['id'] for p in data['programmes']}
ids = [titles[t] for t in pr_ok if t in titles]
if ids:
    lines.append('-- Programmes : seulement si tous leurs entraînements sont validés')
    lines.append('update public.programs p set approved = true where p.id in (' + ', '.join(q(i) for i in sorted(ids)) + ')\n'
                 '  and not exists (select 1 from public.program_days d join public.trainings t on t.id = d.training_id\n'
                 '                  where d.program_id = p.id and not t.approved);')
lines += ['', '-- Ce qui reste à valider après ce script :',
          "select 'exercices' quoi, count(*) from public.exercises where not approved",
          "union all select 'entraînements', count(*) from public.trainings where not approved",
          "union all select 'programmes', count(*) from public.programs where not approved;"]
out = ROOT / 'supabase/content-v5-validation.sql'
out.write_text('\n'.join(lines) + '\n')
print(f'{out} : {len(ex_ok)} exercices, {len(tr_ok)} entraînements, {len(ids)} programmes marqués « oui »')
